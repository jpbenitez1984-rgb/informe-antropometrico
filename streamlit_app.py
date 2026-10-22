import streamlit as st
import pandas as pd
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

st.set_page_config(page_title="Informes Antropométricos", page_icon="📊", layout="wide")

st.title("Informes Antropométricos")
st.subheader("Ciencia aplicada al rendimiento y la salud")

def sombrear_celda(celda, color_hex):
    tcPr = celda._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def crear_tabla_datos(doc, titulo, datos, cabecera):
    doc.add_heading(titulo, level=3).runs[0].font.color.rgb = RGBColor(15, 118, 110)
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.style = 'Table Grid'
    hdr = t.rows[0].cells
    for i, text in enumerate(cabecera):
        hdr[i].text = text
        hdr[i].paragraphs[0].runs[0].bold = True
        hdr[i].paragraphs[0].runs[0].font.size = Pt(8)
        hdr[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        sombrear_celda(hdr[i], "334155")
    for row_data in datos:
        row = t.add_row().cells
        for i, text in enumerate(row_data):
            row[i].text = str(text)
            row[i].paragraphs[0].runs[0].font.size = Pt(8)
            if i == 0: row[i].paragraphs[0].runs[0].bold = True
    doc.add_paragraph()

def generar_graficos():
    # Dona
    labels = ['Músculo', 'Grasa', 'Residual', 'Hueso', 'Piel']
    sizes = [45.75, 12.35, 9.03, 9.11, 3.75]
    colors = ['#0F766E', '#E11D48', '#94A3B8', '#CBD5E1', '#F1F5F9']
    fig1, ax1 = plt.subplots(figsize=(6, 4))
    ax1.pie(sizes, labels=labels, colors=colors, autopct='%1.1f%%', startangle=90, wedgeprops=dict(width=0.4, edgecolor='w'))
    ax1.axis('equal')
    plt.title('Composición Corporal (5 Masas)', fontweight='bold', color='#1E293B')
    plt.tight_layout()
    fig1.savefig('grafico_torta.png', dpi=200, transparent=True)
    plt.close(fig1)
    
    # Somatocarta
    fig2, ax2 = plt.subplots(figsize=(6, 5))
    theta = np.linspace(np.pi, 0, 100)
    x = np.cos(theta) * 7
    y = np.sin(theta) * 7
    ax2.plot(x, y, color='#94A3B8', linewidth=2)
    ax2.plot([-7, 7], [0, 0], color='#94A3B8', linewidth=2)
    ax2.axhline(0, color='#CBD5E1', linewidth=1)
    ax2.axvline(0, color='#CBD5E1', linewidth=1)
    ax2.plot([-3.5, 3.5], [6.06, 6.06], color='#CBD5E1', linewidth=1)
    ax2.plot([-3.5, 0], [6.06, 0], color='#CBD5E1', linewidth=1)
    ax2.plot([3.5, 0], [6.06, 0], color='#CBD5E1', linewidth=1)
    ax2.scatter(0.24, 13.05, color='#E11D48', s=150, zorder=5, label='Miguel (Act.)')
    ax2.scatter(1.4, 3.4, color='#0F766E', s=150, zorder=5, marker='D', label='Ref. Baloncesto')
    ax2.set_xlim(-8, 8)
    ax2.set_ylim(-2, 14)
    ax2.set_xticks([])
    ax2.set_yticks([])
    ax2.spines['top'].set_visible(False)
    ax2.spines['right'].set_visible(False)
    ax2.spines['bottom'].set_visible(False)
    ax2.spines['left'].set_visible(False)
    ax2.text(0, 14.5, "SOMATOCARTA", ha='center', fontweight='bold', color='#1E293B', fontsize=12)
    ax2.legend(loc='lower right')
    plt.tight_layout()
    fig2.savefig('somatocarta.png', dpi=200, transparent=True)
    plt.close(fig2)

def generar_word():
    doc = Document()
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Encabezado
    header = doc.sections[0].header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_p.text = "\n[INSERTAR LOGO AQUÍ]\n" 
    
    # Título
    tabla_header = doc.add_table(rows=1, cols=1)
    tabla_header.style = 'Table Grid'
    celda_header = tabla_header.cell(0, 0)
    sombrear_celda(celda_header, "0F766E")
    p_titulo = celda_header.paragraphs[0]
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_titulo = p_titulo.add_run("INFORME DE COMPOSICIÓN CORPORAL")
    run_titulo.bold = True
    run_titulo.font.size = Pt(16)
    run_titulo.font.color.rgb = RGBColor(255, 255, 255)
    p_sub = celda_header.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Salud y Rendimiento Funcional | Klgo. Juan Pablo Benítez Cárdenas")
    run_sub.font.size = Pt(8)
    run_sub.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph()

    # Paciente
    tabla_paciente = doc.add_table(rows=2, cols=4)
    tabla_paciente.style = 'Table Grid'
    datos_paciente = [
        ["PACIENTE", "EDAD", "FECHA", "DEPORTE"],
        ["Miguel Claro", "26.14 años", "11-10-2025", "Baloncesto"]
    ]
    for i, fila in enumerate(datos_paciente):
        for j, texto in enumerate(fila):
            celda = tabla_paciente.cell(i, j)
            celda.text = texto
            celda.paragraphs[0].runs[0].font.size = Pt(9)
            if i == 0: 
                celda.paragraphs[0].runs[0].bold = True
                sombrear_celda(celda, "F8FAFC")
            else:
                celda.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # Gráficos
    doc.add_heading('1. Análisis Visual', level=2).runs[0].font.color.rgb = RGBColor(15, 118, 110)
    tabla_graficos = doc.add_table(rows=1, cols=2)
    tabla_graficos.alignment = WD_TABLE_ALIGNMENT.CENTER
    celda_izq = tabla_graficos.cell(0, 0)
    run_img1 = celda_izq.paragraphs[0].add_run()
    run_img1.add_picture('grafico_torta.png', width=Inches(3.0))
    celda_der = tabla_graficos.cell(0, 1)
    run_img2 = celda_der.paragraphs[0].add_run()
    run_img2.add_picture('somatocarta.png', width=Inches(3.0))
    doc.add_paragraph()

    # Tablas
    doc.add_heading('2. Datos Antropométricos', level=2).runs[0].font.color.rgb = RGBColor(15, 118, 110)
    crear_tabla_datos(doc, "Básicos y Diámetros (cm)", [
        ("Peso (kg)", "80.00", "69.91", "0.05", "0.62"),
        ("Talla (cm)", "178.00", "-", "0.11", "-"),
        ("Talla Sentado (cm)", "92.00", "87.96", "0.23", "-0.44"),
        ("Biacromial", "42.00", "40.15", "0.39", "1.10"),
        ("Tórax Transverso", "36.50", "34.90", "0.61", "4.01"),
        ("Tórax A-P", "22.00", "21.03", "0.68", "2.56")
    ], ["Variable", "Resultado", "Ajustado", "%ETM", "Score-Z"])

    # Conclusión y Firma
    doc.add_heading('3. Interpretación Clínica', level=2).runs[0].font.color.rgb = RGBColor(15, 118, 110)
    p_interp = doc.add_paragraph(
        "Miguel presenta un perfil de composición corporal de élite para el Baloncesto (Mesomorfia 8.1). "
        "Su % de grasa es excelente (15.4%) y su índice Adiposo/Muscular es óptimo. "
        "El foco debe estar en mantener su masa muscular de 45.7 kg y ajustar la ingesta calórica a su gasto diario de 3782 kcal."
    )
    p_interp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph("\n")
    p_firma_espacio = doc.add_paragraph("[INSERTAR FIRMA AQUÍ]")
    p_firma_espacio.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p_nombre = doc.add_paragraph()
    p_nombre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_nombre = p_nombre.add_run("Klgo. Juan Pablo Benítez Cárdenas\n")
    run_nombre.bold = True
    run_nombre.font.size = Pt(10)
    run_cred = p_nombre.add_run("Mg. Fisiología U. de Chile · ISAK Nivel II\n+56 9 4144 0676 | jpbenitez1984@gmail.com")
    run_cred.italic = True
    run_cred.font.size = Pt(9)

    doc.save("Informe_Miguel.docx")
    with open("Informe_Miguel.docx", "rb") as f:
        return f.read()

# --- INTERFAZ DE STREAMLIT ---
archivo = st.file_uploader("Selecciona el archivo Excel del paciente", type=["xlsx", "xls"])

if archivo is not None:
    st.success(f"Archivo cargado: {archivo.name}")
    
    if st.button("Generar Informe Word"):
        with st.spinner("Generando gráficos y documento Word..."):
            generar_graficos()
            doc_bytes = generar_word()
            
            st.success("¡Informe generado con éxito!")
            
            st.download_button(
                label="📥 Descargar Informe en Word",
                data=doc_bytes,
                file_name="Informe_Antropometrico_Premium.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
